"""Export service — CSV and Excel export for KYC cases."""

from __future__ import annotations

import csv
import io
import datetime as dt
import logging
from typing import Generator

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import Case, Client
from app.security import current_client

logger = logging.getLogger("deepguard.export")

router = APIRouter(prefix="/api/export", tags=["export"])

CHUNK_SIZE = 500


def _generate_csv(query, db: Session) -> Generator[str, None, None]:
    """Stream CSV rows in chunks to avoid loading all data into memory."""
    headers = [
        "id", "applicant", "reference", "country", "doc_type",
        "score", "verdict", "status", "risk_level",
        "created_at", "updated_at", "documents_count",
    ]
    yield ",".join(headers) + "\n"

    offset = 0
    while True:
        cases = query.offset(offset).limit(CHUNK_SIZE).all()
        if not cases:
            break
        for case in cases:
            row = [
                str(case.id),
                case.applicant_name or "",
                case.external_ref or "",
                case.country or "",
                case.documents[0].doc_type.value if case.documents else "",
                str(case.score),
                case.verdict or "",
                case.status.value if case.status else "",
                case.risk_level.value if case.risk_level else "",
                case.created_at.isoformat() if case.created_at else "",
                case.updated_at.isoformat() if case.updated_at else "",
                str(len(case.documents)),
            ]
            yield ",".join(row) + "\n"
        offset += CHUNK_SIZE


@router.get("/cases")
def export_cases(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    status: str | None = None,
    from_date: str | None = None,
    to_date: str | None = None,
    db: Session = Depends(get_db),
    client: Client = Depends(current_client),
):
    """Export KYC cases as CSV or Excel."""
    query = db.query(Case).filter(Case.client_id == client.id)

    if status:
        query = query.filter(Case.status == status)
    if from_date:
        try:
            from_dt = dt.datetime.fromisoformat(from_date)
            query = query.filter(Case.created_at >= from_dt)
        except ValueError:
            pass
    if to_date:
        try:
            to_dt = dt.datetime.fromisoformat(to_date)
            query = query.filter(Case.created_at <= to_dt)
        except ValueError:
            pass

    query = query.order_by(Case.id.desc())

    if format == "csv":
        return StreamingResponse(
            _generate_csv(query, db),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=deepguard-export-{dt.date.today()}.csv"},
        )

    elif format == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "KYC Cases"

            header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)

            headers = [
                "ID", "Applicant", "Reference", "Country", "Doc Type",
                "Score", "Verdict", "Status", "Risk Level",
                "Created At", "Updated At", "Documents Count",
            ]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            row_idx = 2
            offset = 0
            while True:
                cases = query.offset(offset).limit(CHUNK_SIZE).all()
                if not cases:
                    break
                for case in cases:
                    ws.cell(row=row_idx, column=1, value=case.id)
                    ws.cell(row=row_idx, column=2, value=case.applicant_name or "")
                    ws.cell(row=row_idx, column=3, value=case.external_ref or "")
                    ws.cell(row=row_idx, column=4, value=case.country or "")
                    ws.cell(row=row_idx, column=5, value=case.documents[0].doc_type.value if case.documents else "")
                    ws.cell(row=row_idx, column=6, value=case.score)
                    ws.cell(row=row_idx, column=7, value=case.verdict or "")
                    ws.cell(row=row_idx, column=8, value=case.status.value if case.status else "")
                    ws.cell(row=row_idx, column=9, value=case.risk_level.value if case.risk_level else "")
                    ws.cell(row=row_idx, column=10, value=case.created_at.isoformat() if case.created_at else "")
                    ws.cell(row=row_idx, column=11, value=case.updated_at.isoformat() if case.updated_at else "")
                    ws.cell(row=row_idx, column=12, value=len(case.documents))
                    row_idx += 1
                offset += CHUNK_SIZE

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            return StreamingResponse(
                iter([output.getvalue()]),
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename=deepguard-export-{dt.date.today()}.xlsx"},
            )
        except ImportError:
            return {"error": "openpyxl not installed. Run: pip install openpyxl"}
